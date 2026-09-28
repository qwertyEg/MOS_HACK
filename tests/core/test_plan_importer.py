"""Импорт календарного графика: реальные форматы, порча Excel, пустые даты, предупреждения; демо-план."""
import datetime as dt
import io

import openpyxl
import pytest

from core.analytics.timeline import plan_vs_fact
from core.contracts import StageTimeline
from core.plan import importer
from core.plan.sample import ROWS, SAMPLE_PATH

D = dt.date


def _xlsx(rows) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _by_stage(items):
    return {it.stage_id: it for it in items}


# --------------------------------------------------------------------------- пример графика


def test_sample_plan_xlsx_folds_into_eight_stages():
    items, warnings = importer.parse(SAMPLE_PATH.read_bytes(), SAMPLE_PATH.name)
    st = _by_stage(items)
    assert sorted(st) == list(range(1, 9))
    assert (st[1].planned_start, st[1].planned_end) == (D(2026, 2, 2), D(2026, 3, 20))
    # котлован — по детализации, а не по сводной строке «12.3.7.» (она тянется до 31.08 из-за засыпки)
    assert (st[3].planned_start, st[3].planned_end) == (D(2026, 4, 6), D(2026, 6, 5))
    # обратная засыпка пазух ушла в этап 4 и продлила его
    assert "12.3.7./Обратная засыпка" in st[4].work_codes and st[4].planned_end == D(2026, 8, 31)
    assert "12.3.7./Обратная засыпка" not in st[3].work_codes
    assert st[5].planned_end == D(2027, 4, 30)          # наружные стены — часть коробки (этап 5)
    assert st[3].equipment == {"excavator": 2, "dump_truck": 5, "bulldozer": 1, "roller": 1}
    assert st[8].equipment["asphalt_paver"] == 1
    # внутренние работы отброшены с пояснением, других предупреждений на чистом файле нет
    assert len(warnings) == 1 and "внутренних работ: 2" in warnings[0]
    for it in items:
        assert isinstance(it.planned_start, dt.date) and isinstance(it.planned_end, dt.date)


def test_sample_plan_really_contains_excel_date_codes():
    """Файл-пример должен воспроизводить порчу Excel — иначе тест выше ничего не доказывает."""
    ws = openpyxl.load_workbook(SAMPLE_PATH).active
    date_codes = [c.value for c in ws["A"] if isinstance(c.value, dt.datetime)]
    assert dt.datetime(2025, 2, 10) in date_codes          # 10.2.
    assert dt.datetime(2025, 3, 12) in date_codes          # 12.3.
    assert len(ROWS) == ws.max_row - 3


# --------------------------------------------------------------------------- CSV


def test_csv_cp1251_with_semicolons():
    text = ("Код;Наименование работ;Дата начала;Дата окончания;Техника\n"
            "12.3.1.;Устройство котлована;01.04.2026;15.05.2026;Экскаватор 2, самосвал 4\n"
            "12.4.31.;Устройство пирога кровли;01.03.2027;30.04.2027;Башенный кран\n")
    items, warnings = importer.parse(text.encode("cp1251"), "plan.csv")
    st = _by_stage(items)
    assert (st[3].planned_start, st[3].planned_end) == (D(2026, 4, 1), D(2026, 5, 15))
    assert st[3].equipment == {"excavator": 2, "dump_truck": 4}
    assert st[6].planned_start == D(2027, 3, 1)
    assert warnings == []


def test_csv_utf8_bom_commas_english_headers():
    text = "﻿code,name,start,finish\n12.3.1.,Pit excavation,2026-04-01,2026-05-15\n"
    items, warnings = importer.parse(text.encode("utf-8"), "plan.csv")
    assert _by_stage(items)[3].planned_end == D(2026, 5, 15) and warnings == []


def test_excel_date_codes_in_csv_export():
    """CSV, выгруженный из испорченного Excel: вместо «10.2.» стоит «10.02.2025»."""
    text = ("№ п/п;Вид работ;Начало;Окончание\n"
            "10.02.2025;Вынос инженерных систем;02.02.2026;27.02.2026\n"
            "12.03.2025;Устройство подземной части;;\n"
            "12.3.7.;Земляные работы;06.04.2026;29.05.2026\n")
    items, warnings = importer.parse(text.encode("utf-8"), "plan.csv")
    st = _by_stage(items)
    assert st[1].work_codes == ["10.2."] and st[1].planned_start == D(2026, 2, 2)
    assert st[3].planned_end == D(2026, 5, 29)
    assert 10 not in st and 12 not in st            # баг: коды превращались в «этапы 10 и 12»


# --------------------------------------------------------------------------- устойчивость


def test_empty_dates_do_not_crash_and_are_reported():
    """NaT/пустые даты: не падать, не писать 'NaT', предупредить; остальные строки этапа учесть."""
    xlsx = _xlsx([
        ("№ п/п", "Вид работ", "Начало", "Окончание"),
        ("12.3.7.", "Земляные работы", None, None),
        (None, "Разработка грунта", dt.datetime(2026, 4, 6), dt.datetime(2026, 5, 29)),
        (None, "Планировка грунта", None, None),
        ("12.4.31.", "Устройство пирога кровли", "NaT", ""),
    ])
    items, warnings = importer.parse(xlsx, "plan.xlsx")
    st = _by_stage(items)
    assert (st[3].planned_start, st[3].planned_end) == (D(2026, 4, 6), D(2026, 5, 29))
    assert st[6].planned_start is None and st[6].planned_end is None
    assert any("Планировка грунта" in w and "нет дат" in w for w in warnings)
    assert any("этап 6" in w and "без сроков" in w for w in warnings)


def test_unknown_codes_are_warned_not_silenced():
    text = ("Код;Наименование;Начало;Окончание\n"
            "99.9.;Непонятная работа;01.04.2026;15.04.2026\n"
            "12.3.7.9.;Разработка грунта захватка 9;01.04.2026;15.05.2026\n"
            "77.1.;Устройство кровли над паркингом;01.06.2027;30.06.2027\n")
    items, warnings = importer.parse(text.encode("utf-8"), "plan.csv")
    st = _by_stage(items)
    assert any("99.9." in w and "не найден" in w and "не учтена" in w for w in warnings)
    assert st[3].planned_start == D(2026, 4, 1)
    assert any("12.3.7.9." in w and "родительскому коду 12.3.7." in w for w in warnings)
    assert st[6].planned_start == D(2027, 6, 1)                 # по ключевому слову «кровл»
    assert any("77.1." in w and "по названию" in w for w in warnings)


def test_start_after_end_is_swapped_with_warning():
    text = "Код;Наименование;Начало;Окончание\n12.3.1.;Устройство котлована;15.05.2026;01.04.2026\n"
    items, warnings = importer.parse(text.encode("utf-8"), "p.csv")
    it = _by_stage(items)[3]
    assert (it.planned_start, it.planned_end) == (D(2026, 4, 1), D(2026, 5, 15))
    assert any("переставлены" in w for w in warnings)


def test_duration_fills_missing_date_and_bad_date_is_reported():
    text = ("Код;Наименование;Начало;Окончание;Длительность\n"
            "12.3.1.;Устройство котлована;01.04.2026;;10\n"
            "12.4.31.;Кровля;32.13.2026;30.04.2027;\n")
    items, warnings = importer.parse(text.encode("utf-8"), "p.csv")
    st = _by_stage(items)
    assert st[3].planned_end == D(2026, 4, 10)
    assert st[6].planned_end == D(2027, 4, 30) and st[6].planned_start is None
    assert any("32.13.2026" in w for w in warnings)


def test_simple_stage_plan_and_stage_column_with_organizer_codes():
    """Формат ветки api-solution: этап/начало/окончание; подстроки 4.1 и 4.2 сливаются, а не затирают друг друга."""
    text = ("stage_id,start,end\n"
            "3,2026-04-01,2026-05-31\n4.1,2026-06-01,2026-06-20\n4.2,2026-06-15,2026-08-31\n"
            "10.2.,2026-02-01,2026-02-28\n")
    items, warnings = importer.parse(text.encode("utf-8"), "p.csv")
    st = _by_stage(items)
    assert (st[4].planned_start, st[4].planned_end) == (D(2026, 6, 1), D(2026, 8, 31))
    assert st[1].planned_start == D(2026, 2, 1) and 10 not in st
    assert warnings == []


def test_stage_names_without_codes_and_header_not_in_first_row():
    text = ("График производства работ;;\n;;\n"
            "Наименование;Начало;Окончание\n"
            "Котлован;01.04.2026;31.05.2026\n1. Подготовка территории;01.02.2026;20.03.2026\n"
            "Кровля;01.03.2027;30.04.2027\n")
    items, warnings = importer.parse(text.encode("utf-8"), "p.csv")
    assert sorted(_by_stage(items)) == [1, 3, 6] and warnings == []


@pytest.mark.parametrize("data, name, needle", [
    (b"PK\x03\x04garbage", "plan.xlsx", "не читается"),
    ("просто текст без колонок\nещё строка\n".encode("utf-8"), "plan.csv", "заголовков"),
    (b"", "plan.csv", "заголовков"),
    (b"\xd0\xcf\x11\xe0", "plan.xls", ".xls"),
])
def test_bad_files_return_explanation_instead_of_exception(data, name, needle):
    items, warnings = importer.parse(data, name)
    assert items == [] and any(needle in w for w in warnings)


def test_equipment_free_text():
    eq, unknown = importer.parse_equipment("Экскаватор 2; самосвалов ×4, каток, Кран КБ-405, автобетоносмеситель - 3 шт")
    assert eq == {"excavator": 2, "dump_truck": 4, "roller": 1, "concrete_mixer": 3}
    assert unknown == ["Кран КБ-405"]      # «кран» без уточнения неоднозначен, а 405 — не количество
    assert importer.equipment_key("Экскаватор-погрузчик JCB") == "backhoe_loader"
    assert importer.equipment_key("кран-манипулятор") == "crane_manipulator"
    assert importer.equipment_key("Грузовик бортовой") == "truck"


# --------------------------------------------------------------------------- демо-план


def test_demo_plan_is_plausible():
    start = D(2026, 3, 2)
    items = importer.demo_plan(start)
    st = _by_stage(items)
    assert sorted(st) == list(range(1, 9)) and st[1].planned_start == start
    for it in items:
        days = (it.planned_end - it.planned_start).days + 1
        assert 20 <= days <= 300, (it.stage_id, days)
        assert it.equipment and it.work_codes
    # этапы перекрываются, как на реальной стройке, и в целом идут по порядку
    assert st[3].planned_start < st[2].planned_end
    assert st[7].planned_start < st[5].planned_end
    starts = [st[s].planned_start for s in range(1, 6)]
    assert starts == sorted(starts)


def test_demo_plan_is_not_tautological():
    """План под кадры не должен кончаться в дату последнего кадра (иначе ожидается 100 % и вечное «отставание»)."""
    first, last = D(2025, 1, 10), D(2026, 12, 20)          # кадры за два года — дольше типового срока
    items = importer.demo_plan(first, last)
    assert max(it.planned_end for it in items) > last
    pf = plan_vs_fact(items, StageTimeline({}, None, 0.0, []), last)
    assert pf.expected_progress < 0.95
    # короткая серия кадров — нормальные длительности, без растяжения
    short = _by_stage(importer.demo_plan(first, first + dt.timedelta(days=21)))
    assert (short[5].planned_end - short[5].planned_start).days + 1 == 240


def test_demo_plan_stage_subset_starts_at_first_frame():
    items = importer.demo_plan(D(2026, 9, 1), D(2026, 9, 21), stage_ids=[3, 4, 5])
    assert [it.stage_id for it in items] == [3, 4, 5]
    assert items[0].planned_start == D(2026, 9, 1)
